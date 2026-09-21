"""用例库页：测试资产检索、批量操作（建单/停用/删除）、单条编辑与手工新建。"""
from __future__ import annotations

from datetime import datetime

import pandas as pd
import streamlit as st

from excel_export import library_workbook
from methods import METHODS
import store
from styles import (DANGER, PRIMARY, guide_steps, kpi_card, page_header,
                    priority_styler)


@st.cache_data(show_spinner=False)
def _export_library_excel(version: int) -> bytes:
    """用例库 Excel：version 为缓存盐，数据未变时零重建。"""
    return library_workbook(store.list_cases(status="all"))


def render() -> None:
    page_header("用例库", "团队测试资产：检索筛选、圈选建单、编辑维护")
    guide_steps(["筛选或导入用例", "勾选圈选目标用例", "批量新建测试单 / 维护"])

    _asset_overview()
    cases = _toolbar_and_query()
    if not cases:
        st.info("用例库为空。先去「用例生成」采纳 AI 产出，或在下方新建手工用例。")
        _new_case_expander()
        _import_expander()
        return

    selected_ids = _case_table(cases)
    if selected_ids:
        _batch_actions(selected_ids)

    # 维护区三列收拢：新建 / 导入 / 详情编辑 各占一列，页面不再纵向拖长
    c1, c2, c3 = st.columns(3, gap="medium")
    with c1:
        _new_case_expander()
    with c2:
        _import_expander()
    with c3:
        _case_detail(cases)


def _asset_overview() -> None:
    """资产概览：启用 / P0 / AI / 手工 / 停用 五张小指标卡。"""
    all_cases = store.list_cases(status="all")
    active = [c for c in all_cases if c["status"] == "active"]
    stats = [("用例总数（启用）", len(active), PRIMARY),
             ("P0 核心用例", sum(1 for c in active if c["priority"] == "P0"), DANGER),
             ("AI 生成", sum(1 for c in active if c["source"] == "ai"), "#2563EB"),
             ("手工创建", sum(1 for c in active if c["source"] == "manual"), "#64748B"),
             ("已停用", len(all_cases) - len(active), "#94A3B8")]
    cols = st.columns(5)
    for col, (label, value, accent) in zip(cols, stats):
        with col:
            kpi_card(label, value, accent=accent)


def _toolbar_and_query() -> list[dict]:
    """筛选工具栏 → store 查询。"""
    f1, f2, f3, f4, f5 = st.columns([2, 1.4, 1.2, 1, 1])
    keyword = f1.text_input("🔍 关键字（标题/模块/标签）", key="lib_kw")
    sel_method = f2.multiselect("方法", list(METHODS), key="lib_method")
    sel_prio = f3.multiselect("优先级", ["P0", "P1", "P2"], key="lib_prio")
    source = f4.selectbox("来源", ["全部", "AI", "手工"], key="lib_source")
    status = f5.selectbox("状态", ["启用", "停用", "全部"], key="lib_status")

    return store.list_cases(
        keyword=keyword.strip(),
        method=sel_method or None,
        priority=sel_prio or None,
        source={"全部": None, "AI": "ai", "手工": "manual"}[source],
        status={"启用": "active", "停用": "deprecated", "全部": "all"}[status])


def _case_table(cases: list[dict]) -> list[str]:
    """主表格（多选行）→ 返回选中的用例 ID。"""
    rows = [{"ID": c["id"], "模块": c["module"], "标题": c["title"], "方法": c["method"],
             "优先级": c["priority"], "来源": "AI" if c["source"] == "ai" else "手工",
             "标签": c["tags"], "状态": "启用" if c["status"] == "active" else "停用"}
            for c in cases]
    df = pd.DataFrame(rows)
    event = st.dataframe(priority_styler(df), width="stretch", hide_index=True, height=480,
                         on_select="rerun", selection_mode="multi-row", key="lib_table")
    st.caption(f"共 {len(cases)} 条 · 点击行首复选框可多选")
    return [rows[i]["ID"] for i in event.selection.rows]


def _batch_actions(ids: list[str]) -> None:
    a1, a2, a3, a4, a5 = st.columns([2.2, 1.2, 1.2, 1, 2])
    with a1.popover(f"➕ 新建测试单（含选中 {len(ids)} 条）", use_container_width=True):
        name = st.text_input("测试单名称", value=f"回归 {datetime.now():%m-%d %H:%M}", key="mk_run_name")
        env = st.selectbox("环境", ["测试环境", "预发环境", "生产环境"], key="mk_run_env")
        if st.button("创建并前往执行", type="primary", use_container_width=True):
            run_id = store.create_run(name=name, env=env, case_ids=ids)
            st.session_state.enter_run_id = run_id
            st.toast(f"测试单 RUN-{run_id:04d} 已创建")
            from nav import PAGES  # 延迟导入：避免 nav → views → nav 循环
            st.switch_page(PAGES["execution"])
    with a2.popover(f"⏸ 停用（{len(ids)} 条）", use_container_width=True):
        st.caption("停用后不在默认列表与建单圈选中出现，可随时恢复。")
        if st.button("确认停用", use_container_width=True):
            store.set_case_status(ids, "deprecated")
            st.toast(f"已停用 {len(ids)} 条用例")
            st.rerun()
    with a3.popover(f"🗑 删除（{len(ids)} 条）", use_container_width=True):
        st.warning("将同时移除相关执行记录，不可恢复。")
        if st.checkbox("我确认删除", key="del_confirm") and st.button("确认删除", type="primary",
                                                                    use_container_width=True):
            store.delete_cases(ids)
            st.toast(f"已删除 {len(ids)} 条用例")
            st.rerun()
    a4.download_button("⬇️ Excel", data=_export_library_excel(store.version()),
                       file_name="用例库.xlsx",
                       use_container_width=True,
                       help="导出当前用例库全部用例（含停用）")
    with a5:
        from nav import PAGES  # 延迟导入：避免 nav → views → nav 循环
        st.page_link(PAGES["execution"], label="▶️ 前往测试执行", use_container_width=True)


def _case_detail(cases: list[dict]) -> None:
    with st.expander("🔍 用例详情与编辑", expanded=False):
        options = [f'{c["id"]} · {c["title"]}' for c in cases]
        pick = st.selectbox("定位用例", options, key="lib_pick", index=None,
                            placeholder="选择一条用例查看 / 编辑",
                            label_visibility="collapsed")
        if pick is None:
            st.caption("从下拉框选择用例后，可在此查看详情并编辑。")
            return
        case = cases[options.index(pick)]

        d1, d2 = st.columns([1, 1.4], gap="medium")
        with d1:
            meta = (f'<div class="banner"><div class="t">{case["id"]} · {case["title"]}</div>'
                    f'<div class="d">模块：{case["module"]} ｜ 方法：{case["method"]} ｜ '
                    f'来源：{"AI" if case["source"] == "ai" else "手工"}'
                    + (f' ｜ 标签：{case["tags"]}' if case["tags"] else "") + "<br>"
                    f'前置条件：{case["precondition"] or "—"}</div></div>')
            st.markdown(meta, unsafe_allow_html=True)
            st.markdown("**步骤**\n" + case["steps"].replace("\n", "  \n"))
            st.markdown("**预期结果**\n" + (case["expected"] or "—"))
            if case["testpoint_id"]:
                st.caption(f"🔗 追溯测试点：{case['testpoint_id']}")

        with d2:
            with st.form(f"edit_{case['id']}", border=True):
                st.markdown("**编辑用例**")
                t1, t2 = st.columns(2)
                title = t1.text_input("标题", value=case["title"])
                module = t2.text_input("模块", value=case["module"])
                t3, t4, t5 = st.columns(3)
                priority = t3.selectbox("优先级", ["P0", "P1", "P2"],
                                        index=["P0", "P1", "P2"].index(case["priority"]))
                method = t4.selectbox("方法", list(METHODS) + ["手工"],
                                      index=(list(METHODS) + ["手工"]).index(case["method"]))
                tags = t5.text_input("标签（逗号分隔）", value=case["tags"])
                precondition = st.text_area("前置条件", value=case["precondition"])
                steps = st.text_area("步骤（每行一步）", value=case["steps"], height=140)
                expected = st.text_area("预期结果", value=case["expected"], height=90)
                if st.form_submit_button("💾 保存修改", type="primary", use_container_width=True):
                    store.update_case(case["id"], title=title, module=module, priority=priority,
                                      method=method, tags=tags.strip(), precondition=precondition,
                                      steps=steps, expected=expected)
                    st.toast(f"{case['id']} 已保存")
                    st.rerun()


def _new_case_expander() -> None:
    with st.expander("➕ 新建手工用例", expanded=False):
        with st.form("new_case", border=False):
            t1, t2 = st.columns(2)
            title = t1.text_input("标题*", placeholder="如：验证验证码为空时的登录提示")
            module = t2.text_input("模块*", placeholder="如：登录")
            t3, t4, t5 = st.columns(3)
            priority = t3.selectbox("优先级", ["P0", "P1", "P2"], index=1)
            method = t4.selectbox("方法", ["手工"] + list(METHODS))
            tags = t5.text_input("标签（逗号分隔）")
            precondition = st.text_area("前置条件", placeholder="如：已打开登录页")
            steps = st.text_area("步骤（每行一步）", placeholder="1. …\n2. …")
            expected = st.text_area("预期结果", placeholder="如：登录成功，跳转首页")
            submitted = st.form_submit_button("保存新用例", type="primary", use_container_width=True)
        if submitted:
            if not title.strip() or not module.strip():
                st.error("标题与模块为必填项")
            else:
                cid = store.add_case(module=module.strip(), title=title.strip(),
                                     priority=priority, method=method,
                                     precondition=precondition, steps=steps,
                                     expected=expected, tags=tags.strip())
                st.toast(f"手工用例 {cid} 已创建")
                st.rerun()


# ---------------- Excel 批量导入 ----------------

_IMPORT_COLS = {"模块": "module", "标题": "title", "优先级": "priority", "方法": "method",
                "前置条件": "precondition", "步骤": "steps", "预期结果": "expected",
                "标签": "tags"}


def _parse_import(upload) -> list[dict]:
    """解析上传的 Excel：按表头取列，容忍多余列；返回用例 dict 列表。"""
    from openpyxl import load_workbook

    wb = load_workbook(upload, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    colmap = {i: _IMPORT_COLS[h] for i, h in enumerate(headers) if h in _IMPORT_COLS}
    if "title" not in colmap.values():
        raise ValueError("缺少「标题」列——请使用导入模板的表头")

    items = []
    for raw in rows[1:]:
        rec = {name: ("" if raw[i] is None else str(raw[i]).strip())
               for i, name in colmap.items()}
        if rec.get("title") and rec.get("module"):
            rec.setdefault("priority", "P2")
            rec.setdefault("method", "手工")
            items.append(rec)
    return items


def _import_expander() -> None:
    with st.expander("📥 批量导入用例（Excel）", expanded=False):
        st.caption("从 Excel 迁移存量用例：必填列「模块」「标题」，优先级缺省 P2，方法缺省「手工」；"
                   "与用例库导出的 Excel 格式兼容。")
        st.download_button("⬇️ 下载导入模板", data=library_workbook([]),
                           file_name="用例导入模板.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           use_container_width=True)
        upload = st.file_uploader("上传 Excel 文件", type=["xlsx"], key="import_xlsx")
        if upload is None:
            return
        try:
            items = _parse_import(upload)
        except Exception as e:  # noqa: BLE001 - 解析失败统一提示
            st.error(f"解析失败：{e}")
            return
        if not items:
            st.warning("未解析到有效数据行（模块与标题需同时非空）。")
            return

        df = pd.DataFrame([{"导入": True, "模块": it["module"], "标题": it["title"],
                            "优先级": it.get("priority", "P2"), "方法": it.get("method", "手工"),
                            "前置条件": it.get("precondition", ""), "步骤": it.get("steps", ""),
                            "预期结果": it.get("expected", ""), "标签": it.get("tags", "")}
                           for it in items])
        st.caption(f"解析到 **{len(items)}** 条用例，预览确认后入库：")
        ed = st.data_editor(df, key="import_preview", hide_index=True, width="stretch",
                            height=280,
                            column_config={
                                "导入": st.column_config.CheckboxColumn("导入", default=True),
                                "模块": st.column_config.TextColumn("模块", disabled=True),
                                "标题": st.column_config.TextColumn("标题", disabled=True,
                                                                    width="medium"),
                                "优先级": st.column_config.SelectboxColumn(
                                    "优先级", options=["P0", "P1", "P2"], required=True),
                                "前置条件": st.column_config.TextColumn("前置条件", width="large",
                                                                    disabled=True),
                                "步骤": st.column_config.TextColumn("步骤", width="large",
                                                                 disabled=True),
                                "预期结果": st.column_config.TextColumn("预期结果", width="large",
                                                                    disabled=True),
                            })
        n = int(ed["导入"].sum())
        if st.button(f"✅ 导入 {n} 条用例（来源：手工）", type="primary",
                     disabled=n == 0, use_container_width=True, key="import_submit"):
            chosen = [{"module": r["模块"], "title": r["标题"], "priority": r["优先级"],
                       "method": r["方法"] if r["方法"] else "手工",
                       "precondition": r["前置条件"], "steps": r["步骤"],
                       "expected": r["预期结果"], "tags": r["标签"]}
                      for _, r in ed[ed["导入"]].iterrows()]
            ids = store.add_cases(chosen, source="manual")
            st.toast(f"已导入 {len(ids)} 条用例（{ids[0]} ~ {ids[-1]}）")
            st.rerun()
